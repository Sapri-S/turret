"""Build the SkyStrike Fixed Camera plan .docx from its Markdown source.

The Markdown file is the source of truth. Edit that, re-run this, never
hand-edit the .docx -- the next build overwrites it.

    python build_fixed_camera_plan.py
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# Resolves relative to this script so the folder can move without an edit.
ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "SkyStrike_Fixed_Camera_Plan.md"
OUTPUT = ROOT / "SkyStrike_Fixed_Camera_Plan.docx"

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_runs(paragraph, text):
    """Render **bold** and `code` spans as real runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Table: a header row, a separator row, then body rows.
        if stripped.startswith("|") and i + 1 < len(lines) and set(
            lines[i + 1].replace("|", "").replace(" ", "")
        ) <= {"-", ":"} and lines[i + 1].strip().startswith("|"):
            header = split_row(stripped)
            body = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, header):
                cell.text = ""
                add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in body:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = ""
                    add_runs(cell.paragraphs[0], text)
            doc.add_paragraph()
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            heading = doc.add_heading(stripped[2:], level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif stripped.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            add_runs(
                doc.add_paragraph(style="List Number"),
                re.sub(r"^\d+\. ", "", stripped),
            )
        else:
            add_runs(doc.add_paragraph(), stripped)
        i += 1

    doc.save(OUTPUT)
    print("saved", OUTPUT)


if __name__ == "__main__":
    main()
